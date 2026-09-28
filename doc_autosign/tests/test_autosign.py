"""Run from doc_autosign/: python -m unittest discover -s tests -v"""

import hashlib
import io
import re
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

import docx
import pymupdf
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from autosign import demo, stamper  # noqa: E402
from autosign.config import ConfigError, load_config, parse_document_types  # noqa: E402
from autosign.signatures import MissingSignature, SignatureStore, clean_scan  # noqa: E402
from autosign.web import create_app  # noqa: E402
from autosign.workflow import NotAllowed, WorkflowEngine, WorkflowError  # noqa: E402


def images_on_pages(pdf: bytes) -> list[int]:
    doc = pymupdf.open(stream=pdf, filetype="pdf")
    try:
        return [len(p.get_images()) for p in doc]
    finally:
        doc.close()


def pdf_text(pdf: bytes) -> str:
    doc = pymupdf.open(stream=pdf, filetype="pdf")
    try:
        return "".join(p.get_text() for p in doc)
    finally:
        doc.close()


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp)
        self.cfg = load_config(ROOT / "config")
        self.data = self.tmp / "data"
        self.store = SignatureStore(self.cfg, self.data / "signatures")
        for uid in self.cfg.users:
            self.store.save(uid, "full", demo.fake_scan(uid))
        self.store.save("sara.director", "initials", demo.fake_scan("s-i", strokes=2, size=(400, 260)))


class ConfigTests(unittest.TestCase):
    def test_example_config_loads(self):
        cfg = load_config(ROOT / "config")
        self.assertIn("purchase_order", cfg.document_types)
        self.assertEqual([s.id for s in cfg.doc_type("purchase_order").steps], ["prepared", "checked", "approved"])
        self.assertEqual([s.variant for s in cfg.doc_type("purchase_order").step("approved").slots], ["full", "initials"])

    def test_bad_config_is_rejected_with_location(self):
        bad = {"document_types": {"x": {"steps": [{"id": "a", "role": "r", "signature": {"mode": "fixed"}}]}}}
        with self.assertRaisesRegex(ConfigError, r"x\.steps\[1\]\.signature: fixed mode needs"):
            parse_document_types(bad)
        typo = {"document_types": {"x": {"steps": [{"id": "a", "role": "r",
                                                     "signature": {"placeholder": "{{A}}", "widht_mm": 3}}]}}}
        with self.assertRaisesRegex(ConfigError, "unknown signature keys"):
            parse_document_types(typo)


class SignatureTests(Base):
    def test_clean_scan_makes_paper_transparent_and_crops(self):
        png = clean_scan(demo.fake_scan("x"))
        img = Image.open(io.BytesIO(png))
        self.assertEqual(img.mode, "RGBA")
        self.assertEqual(img.getpixel((0, 0))[3], 0)  # corner is paper -> transparent
        self.assertLess(img.width, 900)  # cropped to the ink

    def test_blank_image_is_refused(self):
        blank = io.BytesIO()
        Image.new("RGB", (100, 50), "white").save(blank, "PNG")
        with self.assertRaisesRegex(ValueError, "no ink"):
            clean_scan(blank.getvalue())

    def test_variant_falls_back_to_full(self):
        self.assertEqual(self.store.get("ali.procurement", "initials"), self.store.get("ali.procurement", "full"))

    def test_missing_signature(self):
        empty = SignatureStore(self.cfg, self.tmp / "empty")
        with self.assertRaises(MissingSignature):
            empty.get("ali.procurement")

    def test_path_traversal_in_variant_rejected(self):
        with self.assertRaises(ValueError):
            self.store.save("ali.procurement", "../../evil", demo.fake_scan("x"))


class StamperTests(Base):
    def signer(self):
        return stamper.Signer(name="Sara Khan", designation="Director")

    def test_pdf_text_placeholder_replaced_by_image_and_caption(self):
        dt = self.cfg.doc_type("purchase_order")
        slot = dt.step("approved").slots[0]
        out = stamper.stamp_pdf(demo.purchase_order_pdf(), slot, self.store.get("sara.director"), self.signer())
        text = pdf_text(out)
        self.assertNotIn("{{SIGN_APPROVED}}", text)
        self.assertIn("{{SIGN_PREPARED}}", text)  # other steps untouched
        self.assertIn("Sara Khan", text)
        self.assertEqual(images_on_pages(out), [0, 1])

    def test_occurrence_all_signs_every_page(self):
        slot = self.cfg.doc_type("purchase_order").step("approved").slots[1]
        out = stamper.stamp_pdf(demo.purchase_order_pdf(), slot, self.store.get("sara.director", "initials"),
                                self.signer())
        self.assertEqual(images_on_pages(out), [1, 1])
        self.assertNotIn("{{INITIALS_APPROVER}}", pdf_text(out))

    def test_fixed_position_on_last_page(self):
        slot = self.cfg.doc_type("scanned_form").step("approver").slots[0]
        out = stamper.stamp_pdf(demo.scanned_form_pdf(), slot, self.store.get("sara.director"), self.signer())
        self.assertEqual(images_on_pages(out), [1])

    def test_rotated_pages_get_an_upright_signature_in_the_same_spot(self):
        slot = self.cfg.doc_type("purchase_order").step("approved").slots[0]
        placed = {}
        for rot in (0, 90, 180, 270):
            doc = pymupdf.open()
            page = doc.new_page()
            page.set_rotation(rot)
            page.insert_text(pymupdf.Point(100, 200) * page.derotation_matrix, "OK: {{SIGN_APPROVED}}", rotate=rot)
            out = pymupdf.open(stream=stamper.stamp_pdf(doc.tobytes(), slot, self.store.get("sara.director"),
                                                        self.signer()), filetype="pdf")
            info = out[0].get_image_info(xrefs=True)[0]
            visible = pymupdf.Rect(info["bbox"]) * out[0].rotation_matrix
            placed[rot] = (round(visible.x0), round(visible.y0), round(visible.width) > round(visible.height))
        self.assertEqual(len(set(placed.values())), 1, placed)  # same visible spot, landscape-shaped each time

    def test_fixed_box_outside_page_is_an_error(self):
        slot = self.cfg.doc_type("scanned_form").step("approver").slots[0]
        doc = pymupdf.open()
        doc.new_page().set_rotation(90)  # landscape view: 245 mm down is off the page
        with self.assertRaisesRegex(stamper.PlaceholderNotFound, "outside page"):
            stamper.stamp_pdf(doc.tobytes(), slot, self.store.get("sara.director"), self.signer())

    def test_pdf_form_field(self):
        doc = pymupdf.open()
        page = doc.new_page()
        w = pymupdf.Widget()
        w.field_name, w.field_type, w.rect = "ApproverSign", pymupdf.PDF_WIDGET_TYPE_TEXT, pymupdf.Rect(300, 600, 500, 650)
        page.add_widget(w)
        data = doc.tobytes()
        slot = parse_document_types({"document_types": {"f": {"steps": [
            {"id": "a", "role": "director", "signature": {"mode": "field", "placeholder": "ApproverSign"}}]}}})["f"].steps[0].slots[0]
        out = stamper.stamp_pdf(data, slot, self.store.get("sara.director"), self.signer())
        d2 = pymupdf.open(stream=out, filetype="pdf")
        self.assertEqual(len(list(d2[0].widgets() or [])), 0)
        self.assertEqual(len(d2[0].get_images()), 1)

    def test_missing_placeholder_raises(self):
        slot = self.cfg.doc_type("leave_request").step("hr").slots[0]
        with self.assertRaises(stamper.PlaceholderNotFound):
            stamper.stamp_pdf(demo.purchase_order_pdf(), slot, self.store.get("omar.hr"), self.signer())

    def test_docx_placeholder_in_table(self):
        slot = self.cfg.doc_type("leave_request").step("line_manager").slots[0]
        out = stamper.stamp_docx(demo.leave_request_docx(), slot, self.store.get("priya.finance"),
                                 stamper.Signer(name="Priya Nair"))
        d = docx.Document(io.BytesIO(out))
        cell = d.tables[0].cell(1, 1)
        self.assertNotIn("{{SIGN_MANAGER}}", cell.text)
        self.assertIn("Priya Nair", cell.text)
        self.assertEqual(len(d.inline_shapes), 1)
        self.assertIn("{{SIGN_HR}}", d.tables[0].cell(2, 1).text)

    def test_docx_token_split_across_runs(self):
        d = docx.Document()
        p = d.add_paragraph("Approved: ")
        p.add_run("{{SIGN_")
        p.add_run("MANAGER}}")
        p.add_run(" end")
        buf = io.BytesIO()
        d.save(buf)
        slot = self.cfg.doc_type("leave_request").step("line_manager").slots[0]
        out = stamper.stamp_docx(buf.getvalue(), slot, self.store.get("priya.finance"), stamper.Signer(name="P"))
        d2 = docx.Document(io.BytesIO(out))
        self.assertNotIn("{{", d2.paragraphs[0].text)
        self.assertTrue(d2.paragraphs[0].text.startswith("Approved: "))
        self.assertTrue(d2.paragraphs[0].text.endswith(" end"))
        self.assertEqual(len(d2.inline_shapes), 1)


class WorkflowTests(Base):
    def setUp(self):
        super().setUp()
        self.wf = WorkflowEngine(self.cfg, self.data)

    def submit_po(self):
        return self.wf.submit("purchase_order", "PO-1", "po.pdf", demo.purchase_order_pdf(), "john.staff")

    def test_full_approval_chain_with_audit_page(self):
        req = self.submit_po()
        self.assertEqual(req.current_step, "prepared")
        for step, user in [("prepared", "ali.procurement"), ("checked", "priya.finance"), ("approved", "sara.director")]:
            req = self.wf.approve(req.id, user, step_id=step)
        self.assertEqual(req.status, "approved")
        data, name = self.wf.document(req.id)
        self.assertEqual(name, "po_signed_v3.pdf")
        # p1: initials; p2: three signatures + initials; p3: approval record
        self.assertEqual(images_on_pages(data), [1, 4, 0])
        text = pdf_text(data)
        self.assertIn("Approval record", text)
        self.assertNotIn("{{", text)
        self.assertEqual(self.wf.history(req.id)[-1]["sha256"], hashlib.sha256(data).hexdigest())

    def test_payment_voucher_form_fields(self):
        req = self.wf.submit("payment_voucher", "PV", "pv.pdf", demo.payment_voucher_pdf(), "ali.procurement")
        self.wf.approve(req.id, "priya.finance")
        req = self.wf.approve(req.id, "sara.director")
        data, _ = self.wf.document(req.id)
        doc = pymupdf.open(stream=data, filetype="pdf")
        self.assertEqual(len(list(doc[0].widgets() or [])), 0)
        self.assertEqual(len(doc[0].get_images()), 2)

    def test_seeded_demo(self):
        samples = self.tmp / "samples"
        demo.write_samples(samples)
        made = demo.seed_requests(self.cfg, self.data, samples)
        self.assertEqual(len(made), 4)
        self.assertEqual(sorted(r.status for r in self.wf.list()), ["approved", "pending", "pending", "rejected"])
        self.assertEqual(demo.seed_requests(self.cfg, self.data, samples), [])  # only once

    def test_same_person_cannot_sign_two_steps(self):
        self.cfg.users["ali.procurement"].roles.append("finance")
        req = self.submit_po()
        self.wf.approve(req.id, "ali.procurement")
        with self.assertRaises(NotAllowed):
            self.wf.approve(req.id, "ali.procurement")

    def test_out_of_order_and_wrong_role_refused(self):
        req = self.submit_po()
        with self.assertRaises(NotAllowed):
            self.wf.approve(req.id, "sara.director")  # it's procurement's turn
        with self.assertRaises(WorkflowError):
            self.wf.approve(req.id, "ali.procurement", step_id="checked")  # stale link / wrong step

    def test_reject_stops_flow(self):
        req = self.submit_po()
        req = self.wf.reject(req.id, "ali.procurement", "wrong supplier")
        self.assertEqual(req.status, "rejected")
        with self.assertRaises(WorkflowError):
            self.wf.approve(req.id, "ali.procurement")

    def test_upload_without_placeholders_is_refused(self):
        with self.assertRaisesRegex(WorkflowError, "placeholders missing"):
            self.wf.submit("purchase_order", "x", "x.pdf", demo.scanned_form_pdf(), "john.staff")

    def test_assignable_step_only_named_person(self):
        req = self.wf.submit("leave_request", "Leave", "leave.docx", demo.leave_request_docx(), "john.staff",
                             {"line_manager": "sara.director"})
        with self.assertRaises(NotAllowed):
            self.wf.approve(req.id, "priya.finance")  # has manager role but wasn't picked
        req = self.wf.approve(req.id, "sara.director")
        req = self.wf.approve(req.id, "omar.hr")
        self.assertEqual(req.status, "approved")
        data, name = self.wf.document(req.id)
        self.assertEqual(name, "leave_signed_v2.docx")
        self.assertEqual(len(docx.Document(io.BytesIO(data)).inline_shapes), 2)
        self.assertEqual(len(self.wf.document(req.id, 0)[0]), len(demo.leave_request_docx()))

    def test_assigning_user_without_role_refused(self):
        with self.assertRaisesRegex(WorkflowError, "does not have role"):
            self.wf.submit("leave_request", "Leave", "l.docx", demo.leave_request_docx(), "john.staff",
                           {"line_manager": "omar.hr"})


class WebTests(Base):
    def setUp(self):
        super().setUp()
        self.app = create_app(ROOT / "config", self.data, secret_key="test", api_key="k", base_url="http://x")
        self.app.testing = True
        self.c = self.app.test_client()
        self.engine = self.app.extensions["autosign"]
        self.engine.signatures = self.store

    def csrf(self):
        with self.c.session_transaction() as s:
            s.setdefault("csrf", "t")
            return s["csrf"]

    def login(self, user):
        self.c.get("/login")
        r = self.c.post("/login", data={"user": user, "password": "demo", "csrf": self.csrf()})
        self.assertEqual(r.status_code, 302)

    def test_ui_submit_and_approve(self):
        self.login("john.staff")
        r = self.c.post("/requests/new", data={
            "csrf": self.csrf(), "doc_type": "purchase_order", "title": "PO-7",
            "file": (io.BytesIO(demo.purchase_order_pdf()), "po.pdf")}, content_type="multipart/form-data")
        self.assertEqual(r.status_code, 302, r.data)
        req_id = r.headers["Location"].rsplit("/", 1)[-1]

        # the approver was emailed a link
        outbox = self.engine.notifier.outbox
        self.assertEqual(outbox[-1][0], "ali@example.com")
        link = re.search(r"http://x(/act/\S+)", outbox[-1][2]).group(1)

        # the link page itself does not approve (mail scanners GET links)
        fresh = self.app.test_client()
        self.assertEqual(fresh.get(link).status_code, 200)
        self.assertEqual(self.engine.get(req_id).current_step, "prepared")

        with fresh.session_transaction() as s:
            s["csrf"] = "t"
        r = fresh.post(link, data={"csrf": "t", "action": "approve"})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(self.engine.get(req_id).current_step, "checked")
        # reusing the link does nothing
        fresh.post(link, data={"csrf": "t", "action": "approve"})
        self.assertEqual(self.engine.get(req_id).version, 1)

        # logged-in approval
        self.c.post("/logout", data={"csrf": self.csrf()})
        self.login("priya.finance")
        r = self.c.post(f"/requests/{req_id}/act", data={"csrf": self.csrf(), "action": "approve", "step": "checked"})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(self.engine.get(req_id).current_step, "approved")

    def test_post_without_csrf_rejected(self):
        self.login("john.staff")
        self.assertEqual(self.c.post("/requests/new", data={}).status_code, 400)

    def test_outsider_cannot_view(self):
        req = self.engine.submit("purchase_order", "PO", "po.pdf", demo.purchase_order_pdf(), "ali.procurement")
        self.login("omar.hr")
        self.assertEqual(self.c.get(f"/requests/{req.id}").status_code, 403)

    def test_login_page_lists_demo_accounts(self):
        page = self.c.get("/login").get_data(as_text=True)
        self.assertIn("Demo accounts", page)
        for uid in self.cfg.users:
            self.assertIn(uid, page)

    def test_admin_page_admin_only(self):
        self.engine.submit("purchase_order", "PO-A", "po.pdf", demo.purchase_order_pdf(), "john.staff")
        self.login("priya.finance")
        self.assertEqual(self.c.get("/admin").status_code, 403)
        self.c.post("/logout", data={"csrf": self.csrf()})
        self.login("sara.director")
        page = self.c.get("/admin").get_data(as_text=True)
        self.assertIn("PO-A", page)
        self.assertIn("Users and signatures", page)
        self.assertIn("Payment Voucher", page)

    def test_admin_cancel_and_signature_upload(self):
        req = self.engine.submit("purchase_order", "PO-C", "po.pdf", demo.purchase_order_pdf(), "john.staff")
        self.login("sara.director")
        self.c.post(f"/admin/requests/{req.id}/cancel", data={"csrf": self.csrf()})
        self.assertEqual(self.engine.get(req.id).status, "cancelled")
        r = self.c.post("/admin/users/omar.hr/signature", content_type="multipart/form-data", data={
            "csrf": self.csrf(), "variant": "initials", "file": (io.BytesIO(demo.fake_scan("o")), "o.jpg")})
        self.assertEqual(r.status_code, 302)
        self.assertIn("initials", self.engine.signatures.variants("omar.hr"))
        self.assertEqual(self.c.get("/admin/users/omar.hr/signature/initials.png").status_code, 200)

    def test_api_flow(self):
        h = {"X-API-Key": "k"}
        self.assertEqual(self.c.post("/api/requests", headers={"X-API-Key": "bad"}).status_code, 401)
        r = self.c.post("/api/requests", headers=h, content_type="multipart/form-data", data={
            "doc_type": "leave_request", "submitter": "john.staff", "assign_line_manager": "priya.finance",
            "file": (io.BytesIO(demo.leave_request_docx()), "leave.docx")})
        self.assertEqual(r.status_code, 201, r.json)
        rid = r.json["id"]
        r = self.c.post(f"/api/requests/{rid}/approve", headers=h, json={"user": "omar.hr"})
        self.assertEqual(r.status_code, 403)
        r = self.c.post(f"/api/requests/{rid}/approve", headers=h, json={"user": "priya.finance", "step": "line_manager"})
        self.assertEqual(r.json["current_step"], "hr")

    def test_api_stateless_sign(self):
        r = self.c.post("/api/sign", headers={"X-API-Key": "k"}, content_type="multipart/form-data", data={
            "doc_type": "purchase_order", "step": "prepared", "user": "ali.procurement",
            "file": (io.BytesIO(demo.purchase_order_pdf()), "po.pdf")})
        self.assertEqual(r.status_code, 200)
        self.assertNotIn("{{SIGN_PREPARED}}", pdf_text(r.data))


if __name__ == "__main__":
    unittest.main()
