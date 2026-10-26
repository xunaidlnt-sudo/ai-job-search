"""Puts a signature image (and an optional name/date caption) onto a document.

PDF: the placeholder can be a text token such as {{SIGN_APPROVER}}, a form
field name, or a fixed box in millimetres (for scanned PDFs with no text
layer). DOCX: text tokens only - the signature goes inline where the token was.
"""

from __future__ import annotations

import copy
import io
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Callable

import pymupdf
from PIL import Image

from .config import MM, DocumentType, SignatureSlot, Step


class PlaceholderNotFound(LookupError):
    pass


class UnsupportedDocument(ValueError):
    pass


@dataclass
class Signer:
    name: str
    designation: str = ""
    email: str = ""
    step_label: str = ""
    when: datetime | None = None


class _Blank(dict):
    def __missing__(self, key):
        return ""


def caption_text(slot: SignatureSlot, signer: Signer) -> str:
    when = signer.when or datetime.now()
    values = _Blank(
        name=signer.name,
        designation=signer.designation,
        email=signer.email,
        label=signer.step_label,
        date=when.strftime(slot.date_format),
    )
    lines = [line for line in slot.caption.format_map(values).splitlines() if line.strip()]
    return "\n".join(lines)


def doc_kind(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return "pdf"
    if ext == ".docx":
        return "docx"
    raise UnsupportedDocument(f"only .pdf and .docx are supported, got '{ext or filename}'")


def _pick(items: list, occurrence) -> list:
    if not items:
        return []
    if occurrence == "all":
        return items
    if occurrence == "first":
        return items[:1]
    if occurrence == "last":
        return items[-1:]
    idx = int(occurrence) - 1
    return items[idx:idx + 1]


# ---------------------------------------------------------------- PDF

def _pdf_targets(doc: pymupdf.Document, slot: SignatureSlot) -> list[tuple[int, pymupdf.Rect, int | None]]:
    """(page index, placeholder rect, widget xref or None) for each spot to sign."""
    found: list[tuple[int, pymupdf.Rect, int | None]] = []
    if slot.mode == "text":
        for page in doc:
            for rect in page.search_for(slot.placeholder):
                found.append((page.number, rect, None))
    elif slot.mode == "field":
        for page in doc:
            for w in page.widgets() or []:
                if w.field_name == slot.placeholder:
                    found.append((page.number, pymupdf.Rect(w.rect), w.xref))
    else:  # fixed
        pno = slot.page - 1 if slot.page > 0 else doc.page_count + slot.page
        if not 0 <= pno < doc.page_count:
            raise PlaceholderNotFound(f"page {slot.page} does not exist (document has {doc.page_count})")
        x, y, w, h = (v * MM for v in slot.rect_mm)
        # rect_mm is measured on the page as the reader sees it; store it in
        # the page's own (unrotated) coordinates like the other modes.
        visible = pymupdf.Rect(x, y, x + w, y + h)
        if not doc[pno].rect.contains(visible):
            raise PlaceholderNotFound(f"fixed box {list(slot.rect_mm)} mm lies outside page {pno + 1}")
        found.append((pno, visible * doc[pno].derotation_matrix, None))
    return _pick(found, slot.occurrence)


def _box(slot: SignatureSlot, ph: pymupdf.Rect) -> pymupdf.Rect:
    if slot.mode in ("fixed", "field"):
        # The box is the field / fixed rectangle itself.
        box = pymupdf.Rect(ph)
    else:
        w, h = slot.width_mm * MM, slot.height_mm * MM
        cx, cy = (ph.x0 + ph.x1) / 2, (ph.y0 + ph.y1) / 2
        x0, y0 = {
            "over": (cx - w / 2, cy - h / 2),
            "above": (cx - w / 2, ph.y0 - h),
            "below": (cx - w / 2, ph.y1),
            "right": (ph.x1, cy - h / 2),
            "left": (ph.x0 - w, cy - h / 2),
        }[slot.anchor]
        box = pymupdf.Rect(x0, y0, x0 + w, y0 + h)
    dx, dy = slot.offset_mm
    return box + (dx * MM, dy * MM, dx * MM, dy * MM)


def _image_rect(box: pymupdf.Rect, png: bytes) -> pymupdf.Rect:
    """Largest rect with the image's aspect ratio that fits in `box`, centred."""
    with Image.open(io.BytesIO(png)) as img:
        ratio = img.width / img.height
    w = min(box.width, box.height * ratio)
    h = w / ratio
    x0 = box.x0 + (box.width - w) / 2
    y0 = box.y0 + (box.height - h) / 2
    return pymupdf.Rect(x0, y0, x0 + w, y0 + h)


def stamp_pdf(data: bytes, slot: SignatureSlot, signature_png: bytes, signer: Signer) -> bytes:
    doc = pymupdf.open(stream=data, filetype="pdf")
    try:
        targets = _pdf_targets(doc, slot)
        if not targets:
            if slot.required:
                raise PlaceholderNotFound(f"placeholder '{slot.placeholder}' not found in the PDF")
            return data

        # Clear placeholders first; redaction must run before the image goes on
        # or it would clip anything overlapping the token's rectangle.
        if slot.remove_placeholder:
            by_page: dict[int, list] = {}
            for pno, rect, xref in targets:
                by_page.setdefault(pno, []).append((rect, xref))
            for pno, spots in by_page.items():
                page = doc[pno]
                redact = False
                for rect, xref in spots:
                    if xref is not None:
                        page.delete_widget(page.load_widget(xref))
                    elif slot.mode == "text":
                        page.add_redact_annot(rect, fill=False)
                        redact = True
                if redact:
                    page.apply_redactions(
                        images=pymupdf.PDF_REDACT_IMAGE_NONE,
                        graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
                    )

        caption = caption_text(slot, signer)
        for pno, rect, _ in targets:
            page = doc[pno]
            # Lay out in the rotated view (what the reader sees) so "above",
            # width/height and the caption line up on rotated scans too, then
            # map back to page coordinates for drawing.
            to_page = page.derotation_matrix
            box = _box(slot, rect * page.rotation_matrix)
            img_rect = _image_rect(box, signature_png)
            page.insert_image(img_rect * to_page, stream=signature_png, keep_proportion=True,
                              overlay=True, rotate=page.rotation)
            if caption:
                _pdf_caption(page, img_rect, box, caption, slot.caption_font_size)
        return doc.tobytes(garbage=3, deflate=True)
    finally:
        doc.close()


def _pdf_caption(page: pymupdf.Page, img: pymupdf.Rect, box: pymupdf.Rect, text: str, size: float) -> None:
    lines = text.count("\n") + 1
    height = lines * size * 1.25 + 2
    extra = 15 * MM  # let long names run a little wider than the signature box
    rect = pymupdf.Rect(box.x0 - extra, img.y1 + 1, box.x1 + extra, img.y1 + 1 + height)
    page.insert_textbox(rect * page.derotation_matrix, text, fontsize=size, fontname="helv",
                        align=pymupdf.TEXT_ALIGN_CENTER, rotate=page.rotation)


def append_audit_page(data: bytes, title: str, lines: list[str]) -> bytes:
    """Adds a plain approval certificate page listing who signed and when."""
    doc = pymupdf.open(stream=data, filetype="pdf")
    try:
        last = doc[-1].rect if doc.page_count else pymupdf.paper_rect("a4")
        page = doc.new_page(width=last.width, height=last.height)
        margin = 20 * MM
        page.insert_text((margin, margin + 10), title, fontsize=14, fontname="hebo")
        body = pymupdf.Rect(margin, margin + 25, last.width - margin, last.height - margin)
        page.insert_textbox(body, "\n".join(lines), fontsize=9, fontname="helv")
        return doc.tobytes(garbage=3, deflate=True)
    finally:
        doc.close()


# ---------------------------------------------------------------- DOCX

def _docx_paragraphs(document):
    """Every paragraph in the body, tables (nested), headers and footers."""

    def walk(container):
        for p in container.paragraphs:
            yield p
        for table in getattr(container, "tables", []):
            for row in table.rows:
                for cell in row.cells:
                    yield from walk(cell)

    yield from walk(document)
    seen = set()
    for section in document.sections:
        for part in (section.header, section.footer, section.first_page_header,
                     section.first_page_footer, section.even_page_header, section.even_page_footer):
            if part.is_linked_to_previous:
                continue
            key = id(part._element)
            if key in seen:
                continue
            seen.add(key)
            yield from walk(part)


def _isolate_token(paragraph, token: str, nth: int):
    """Split runs so the nth token in the paragraph sits alone in one run and
    return that run. Word often splits a token across runs (spell-check,
    formatting edits); in that case the paragraph's runs are merged first,
    which keeps the first run's formatting for the whole paragraph."""
    from docx.text.run import Run

    for attempt in range(2):
        seen = 0
        for run in paragraph.runs:
            count = run.text.count(token)
            if seen + count > nth:
                text = run.text
                pos = -1
                for _ in range(nth - seen + 1):
                    pos = text.index(token, pos + 1)
                before, after = text[:pos], text[pos + len(token):]
                run.text = before
                token_r = copy.deepcopy(run._r)
                run._r.addnext(token_r)
                token_run = Run(token_r, paragraph)
                token_run.text = ""
                if after:
                    after_r = copy.deepcopy(run._r)
                    token_r.addnext(after_r)
                    Run(after_r, paragraph).text = after
                return token_run
            seen += count
        if attempt == 0:
            runs = paragraph.runs
            if not runs:
                break
            merged = "".join(r.text for r in runs)
            runs[0].text = merged
            for r in runs[1:]:
                r._r.getparent().remove(r._r)
    raise PlaceholderNotFound(f"could not isolate '{token}' in paragraph")


def stamp_docx(data: bytes, slot: SignatureSlot, signature_png: bytes, signer: Signer) -> bytes:
    import docx
    from docx.shared import Mm, Pt

    if slot.mode != "text":
        raise UnsupportedDocument("Word documents support text placeholders only (mode: text)")
    document = docx.Document(io.BytesIO(data))
    token = slot.placeholder
    spots = []
    for p in _docx_paragraphs(document):
        for i in range(p.text.count(token)):
            spots.append((p, i))
    spots = _pick(spots, slot.occurrence)
    if not spots:
        if slot.required:
            raise PlaceholderNotFound(f"placeholder '{token}' not found in the Word document")
        return data

    with Image.open(io.BytesIO(signature_png)) as img:
        ratio = img.width / img.height
    width_mm = min(slot.width_mm, slot.height_mm * ratio)
    caption = caption_text(slot, signer)

    # Replace from the end so earlier indexes in the same paragraph stay valid.
    for paragraph, nth in reversed(spots):
        run = _isolate_token(paragraph, token, nth)
        run.add_picture(io.BytesIO(signature_png), width=Mm(width_mm), height=Mm(width_mm / ratio))
        if caption:
            run.font.size = Pt(slot.caption_font_size)
            for line in caption.split("\n"):
                run.add_break()
                run.add_text(line)

    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def docx_to_pdf(data: bytes) -> bytes:
    """Converts with LibreOffice in headless mode. Raises if it is not installed."""
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        raise RuntimeError("LibreOffice (soffice) is needed to convert Word files to PDF")
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "document.docx"
        src.write_bytes(data)
        # A private profile lets this run while someone has LibreOffice open.
        profile = (Path(tmp) / "profile").as_uri()
        proc = subprocess.run(
            [exe, f"-env:UserInstallation={profile}", "--headless", "--convert-to", "pdf",
             "--outdir", tmp, str(src)],
            capture_output=True, text=True, timeout=180,
        )
        out = Path(tmp) / "document.pdf"
        if not out.is_file():
            # soffice exits 0 even when it fails, e.g. when Writer is not installed.
            raise RuntimeError("LibreOffice could not convert the file (is libreoffice-writer installed?): "
                               + (proc.stderr or proc.stdout).strip()[-300:])
        return out.read_bytes()


# ---------------------------------------------------------------- dispatch

def stamp(data: bytes, filename: str, slot: SignatureSlot, signature_png: bytes, signer: Signer) -> bytes:
    kind = doc_kind(filename)
    if kind == "pdf":
        return stamp_pdf(data, slot, signature_png, signer)
    return stamp_docx(data, slot, signature_png, signer)


def sign_step(data: bytes, filename: str, step: Step, signature_for: Callable[[str], bytes], signer: Signer) -> bytes:
    """Apply every signature slot of one approval step. `signature_for(variant)`
    returns the approver's PNG for that variant."""
    for slot in step.slots:
        data = stamp(data, filename, slot, signature_for(slot.variant), signer)
    return data


def check_placeholders(data: bytes, filename: str, doc_type: DocumentType) -> dict[str, list[int]]:
    """For each step, how many signing spots each of its slots would hit in
    this file. Used to catch template mistakes when a document is uploaded,
    not at the last approval."""
    kind = doc_kind(filename)
    result: dict[str, list[int]] = {}
    if kind == "pdf":
        doc = pymupdf.open(stream=data, filetype="pdf")
        try:
            for step in doc_type.steps:
                counts = []
                for slot in step.slots:
                    try:
                        counts.append(len(_pdf_targets(doc, slot)))
                    except PlaceholderNotFound:
                        counts.append(0)
                result[step.id] = counts
        finally:
            doc.close()
    else:
        import docx

        document = docx.Document(io.BytesIO(data))
        text = [p.text for p in _docx_paragraphs(document)]
        for step in doc_type.steps:
            counts = []
            for slot in step.slots:
                if slot.mode != "text":
                    counts.append(0)
                    continue
                total = sum(t.count(slot.placeholder) for t in text)
                counts.append(len(_pick(list(range(total)), slot.occurrence)))
            result[step.id] = counts
    return result


def missing_placeholders(counts: dict[str, list[int]], doc_type: DocumentType) -> list[str]:
    """Human-readable list of required slots with nothing to sign."""
    missing = []
    for step in doc_type.steps:
        for slot, n in zip(step.slots, counts.get(step.id, [])):
            if slot.required and n == 0:
                missing.append(f"{step.label} ({slot.placeholder or 'fixed box'})")
    return missing
