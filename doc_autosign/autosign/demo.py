"""Sample documents and fake scanned signatures so the whole flow can be tried
without real files. Nothing here is needed in production."""

from __future__ import annotations

import io
import math
import random
from pathlib import Path

import pymupdf
from PIL import Image, ImageDraw, ImageFilter

from .config import AppConfig
from .signatures import SignatureStore


def fake_scan(seed: str, strokes: int = 5, size=(900, 300)) -> bytes:
    """A pen-like squiggle in blue ink on slightly grey paper, saved as JPEG -
    close enough to a phone scan to exercise the background cleanup."""
    rnd = random.Random(seed)
    img = Image.new("RGB", size, (238, 236, 230))
    draw = ImageDraw.Draw(img)
    w, h = size
    x = 40
    for _ in range(strokes):
        pts = []
        span = rnd.randint(w // (strokes + 2), w // strokes)
        amp, freq, base = rnd.uniform(30, 80), rnd.uniform(2, 5), h / 2 + rnd.uniform(-30, 30)
        for i in range(80):
            t = i / 79
            pts.append((x + t * span, base - amp * math.sin(t * freq * math.pi) * (1 - t * 0.4)))
        draw.line(pts, fill=(20, 40, 140), width=rnd.randint(5, 8), joint="curve")
        x += span * rnd.uniform(0.7, 0.95)
    draw.line([(60, h * 0.78), (min(x, w - 40), h * 0.72)], fill=(20, 40, 140), width=5)
    img = img.filter(ImageFilter.GaussianBlur(1.2))
    out = io.BytesIO()
    img.save(out, "JPEG", quality=85)
    return out.getvalue()


def _letterhead(page: pymupdf.Page, title: str, ref: str) -> None:
    page.draw_rect(pymupdf.Rect(0, 0, page.rect.width, 8), color=None, fill=(0.12, 0.25, 0.6))
    page.insert_text((50, 42), "ACME Engineering & Contracting LLC", fontsize=9, color=(0.35, 0.35, 0.35))
    page.insert_text((50, 70), title, fontsize=18, fontname="hebo")
    page.insert_text((50, 88), ref, fontsize=10)
    page.draw_line((50, 98), (545, 98), width=0.5, color=(0.7, 0.7, 0.7))


def _table(page: pymupdf.Page, y: float, cols: list[float], header: tuple, rows: list[tuple]) -> float:
    for x, cell in zip(cols, header):
        page.insert_text((x, y), cell, fontsize=9, fontname="hebo")
    y += 6
    page.draw_line((50, y), (545, y), width=0.4)
    y += 14
    for r in rows:
        for x, cell in zip(cols, r):
            page.insert_text((x, y), cell, fontsize=10)
        y += 16
    page.draw_line((50, y - 8), (545, y - 8), width=0.4)
    return y


def purchase_order_pdf(po: str = "PO-2291", supplier: str = "Gulf Cable Trading LLC",
                       project: str = "132kV Substation Extension", rows: list[tuple] | None = None) -> bytes:
    rows = rows or [("1", "Cable tray 300mm HDG", "120 m", "150.00", "18,000.00"),
                    ("2", "Tray bends 90 deg", "24 no", "120.00", "2,880.00"),
                    ("3", "Support brackets", "200 no", "23.00", "4,600.00")]
    total = sum(float(r[4].replace(",", "")) for r in rows)
    doc = pymupdf.open()
    for n in (1, 2):
        page = doc.new_page(width=595, height=842)
        _letterhead(page, "PURCHASE ORDER", f"{po}  |  Page {n} of 2")
        if n == 1:
            y = 125
            for line in [f"Supplier: {supplier}", f"Project: {project}",
                         "Delivery: Site store, within 21 days of this order"]:
                page.insert_text((50, y), line, fontsize=11)
                y += 18
            y = _table(page, y + 20, [50, 80, 300, 380, 470], ("#", "Description", "Qty", "Rate", "Amount"), rows)
            page.insert_text((380, y + 8), f"Total (AED): {total:,.2f}", fontsize=11, fontname="hebo")
        else:
            y = 125
            for line in ["Terms and conditions",
                         "1. Payment 30 days from delivery against a valid tax invoice.",
                         "2. Prices exclude VAT.",
                         "3. Material must match the approved submittal; rejected items are returned at supplier cost.",
                         "4. This order is valid only when signed by all three signatories below."]:
                page.insert_text((50, y), line, fontsize=10, fontname="hebo" if line.startswith("Terms") else "helv")
                y += 17
            y = 610
            for x, label, token in ((50, "Prepared by", "{{SIGN_PREPARED}}"),
                                    (230, "Checked by", "{{SIGN_CHECKED}}"),
                                    (410, "Approved by", "{{SIGN_APPROVED}}")):
                page.insert_text((x + 10, y), token, fontsize=9, color=(0.55, 0.55, 0.55))
                page.draw_line((x, y + 50), (x + 150, y + 50), width=0.6)
                page.insert_text((x, y + 64), label, fontsize=10)
        page.insert_text((440, 815), "Initials: {{INITIALS_APPROVER}}", fontsize=8)
    data = doc.tobytes()
    doc.close()
    return data


def payment_voucher_pdf(pv: str = "PV-0418") -> bytes:
    """A fillable PDF whose signature boxes are form fields, not text tokens."""
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    _letterhead(page, "PAYMENT VOUCHER", pv)
    y = 125
    for line in ["Payee: Al Noor Scaffolding Services", "Invoice: INV-7731 dated 14-Sep-2026",
                 "Cost code: 4410-SCF  |  Project: 132kV Substation Extension"]:
        page.insert_text((50, y), line, fontsize=11)
        y += 18
    y = _table(page, y + 20, [50, 80, 380, 470], ("#", "Description", "", "Amount"),
               [("1", "Scaffolding hire, August 2026", "", "11,200.00"),
                ("2", "Erection and dismantling", "", "3,450.00")])
    page.insert_text((380, y + 8), "Net payable (AED): 14,650.00", fontsize=11, fontname="hebo")
    for x, label, name in ((50, "Verified by (Accounts)", "Sig_Accounts"), (320, "Authorised by", "Sig_Director")):
        rect = pymupdf.Rect(x, 560, x + 220, 620)
        w = pymupdf.Widget()
        w.field_name, w.field_type, w.rect = name, pymupdf.PDF_WIDGET_TYPE_TEXT, rect
        w.border_color, w.border_width = (0.6, 0.6, 0.6), 0.5
        page.add_widget(w)
        page.insert_text((x, 650), label, fontsize=10)
    data = doc.tobytes()
    doc.close()
    return data


def scanned_form_pdf() -> bytes:
    """Stands in for a scanned page: the signature goes in a fixed box because
    a real scan has no text to search for."""
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    _letterhead(page, "SITE ACCESS PERMIT", "SAP-1150  |  scanned copy")
    y = 125
    for line in ["Visitor: Delta Testing Services (3 persons)", "Area: GIS building, 132kV switchgear room",
                 "Date: 05-Oct-2026, 07:00 - 17:00", "Purpose: Primary injection testing, bays 3 and 4"]:
        page.insert_text((50, y), line, fontsize=11)
        y += 18
    page.draw_rect(pymupdf.Rect(130, 245, 180, 263) * (72 / 25.4), color=(0.7, 0.7, 0.7), width=0.5)
    page.insert_text((130 * 72 / 25.4, 270 * 72 / 25.4), "Approver signature", fontsize=9)
    data = doc.tobytes()
    doc.close()
    return data


def leave_request_docx(employee: str = "John Mathew, Site Engineer",
                       dates: str = "12 Oct 2026 to 16 Oct 2026 (5 working days), annual leave") -> bytes:
    import docx
    from docx.shared import Pt

    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "ACME Engineering & Contracting LLC - HR Form 07"
    d.add_heading("Leave Request", level=1)
    d.add_paragraph(f"Employee: {employee}")
    d.add_paragraph(f"Dates: {dates}.")
    d.add_paragraph("Handover: site activities handed to the deputy site engineer for the period.")
    d.add_paragraph()
    table = d.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    for i, (label, token) in enumerate([("Approval", ""), ("Line manager", "{{SIGN_MANAGER}}"),
                                        ("HR", "{{SIGN_HR}}")]):
        table.cell(i, 0).text = label
        table.cell(i, 1).text = token
    table.cell(0, 0).paragraphs[0].runs[0].font.bold = True
    for row in table.rows[1:]:
        row.cells[1].paragraphs[0].runs[0].font.size = Pt(9)
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


SAMPLE_FILES = {
    "purchase_order_PO-2291.pdf": lambda: purchase_order_pdf(),
    "purchase_order_PO-2288.pdf": lambda: purchase_order_pdf(
        "PO-2288", "Emirates Fasteners Trading", "Warehouse Racking Upgrade",
        [("1", "Anchor bolts M16 x 150", "400 no", "6.50", "2,600.00"),
         ("2", "Chemical anchor cartridges", "40 no", "48.00", "1,920.00")]),
    "payment_voucher_PV-0418.pdf": lambda: payment_voucher_pdf(),
    "leave_request_john.docx": lambda: leave_request_docx(),
    "scanned_site_access_permit.pdf": lambda: scanned_form_pdf(),
    "signature_scan_example.jpg": lambda: fake_scan("example"),
}


def write_samples(folder: Path, overwrite: bool = False) -> list[Path]:
    folder.mkdir(parents=True, exist_ok=True)
    written = []
    for name, make in SAMPLE_FILES.items():
        path = folder / name
        if overwrite or not path.exists():
            path.write_bytes(make())
            written.append(path)
    return written


def give_everyone_signatures(cfg: AppConfig, data_dir: Path) -> None:
    store = SignatureStore(cfg, data_dir / "signatures")
    for uid in cfg.users:
        if "full" not in store.variants(uid):
            store.save(uid, "full", fake_scan(uid))
        if "initials" not in store.variants(uid):
            store.save(uid, "initials", fake_scan(uid + "-i", strokes=2, size=(400, 260)))


def seed_requests(cfg: AppConfig, data_dir: Path, samples: Path) -> list[str]:
    """A few requests in different states so the dashboards are not empty.
    Does nothing if the database already has requests."""
    from .workflow import WorkflowEngine

    wf = WorkflowEngine(cfg, data_dir)
    if wf.list():
        return []
    read = lambda name: (samples / name).read_bytes()  # noqa: E731
    made = []

    r = wf.submit("purchase_order", "PO-2291 cable trays", "purchase_order_PO-2291.pdf",
                  read("purchase_order_PO-2291.pdf"), "john.staff")
    wf.approve(r.id, "ali.procurement", "Quotes compared, lowest compliant bidder.")
    made.append(f"{r.title}: prepared, waiting for Finance (priya.finance)")

    r = wf.submit("payment_voucher", "PV-0418 scaffolding August", "payment_voucher_PV-0418.pdf",
                  read("payment_voucher_PV-0418.pdf"), "ali.procurement")
    wf.approve(r.id, "priya.finance")
    wf.approve(r.id, "sara.director")
    made.append(f"{r.title}: fully approved")

    r = wf.submit("leave_request", "Leave - John Mathew, 12-16 Oct", "leave_request_john.docx",
                  read("leave_request_john.docx"), "john.staff", {"line_manager": "sara.director"})
    made.append(f"{r.title}: waiting for line manager (sara.director)")

    r = wf.submit("purchase_order", "PO-2288 anchor bolts", "purchase_order_PO-2288.pdf",
                  read("purchase_order_PO-2288.pdf"), "john.staff")
    wf.approve(r.id, "ali.procurement")
    wf.reject(r.id, "priya.finance", "Supplier quote expired on 20-Sep. Please re-quote.")
    made.append(f"{r.title}: rejected by Finance")
    return made
